"""Read an uploaded document, find the links inside it, and check quotes.

The safety idea in one line: the AI must quote the document word for word,
and this module proves the quote is there before anything is shown.
"""
from __future__ import annotations

import io
import json
import os
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf

ROOT = Path(__file__).parent
DATA = ROOT / "data"
SAMPLE_PDF = DATA / "sample.pdf"

MIN_QUOTE_CHARS = 20  # after whitespace is removed; blocks trivial "quotes"
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_MB", "") or 15) * 1024 * 1024
MAX_PDF_PAGES = 200
MAX_TEXT_CHARS = 300_000  # about 75k tokens


class DocumentError(Exception):
    """A problem with the file that can be explained to the person in plain words."""


@dataclass
class Page:
    number: int  # number used in [PAGE n] markers and quotes (unique within a Source)
    text: str
    printed: str | None = None  # number printed on the page, if any
    origin: str | None = None  # title of the linked page this came from
    url: str | None = None
    local: int | None = None  # page or part number inside its own origin
    official: bool | None = None


@dataclass
class LinkRef:
    url: str
    page: int | None = None  # page of the document the link was found on


@dataclass
class Source:
    """One body of text the bot may answer from."""

    name: str
    kind: str  # "main" or "linked"
    pages: list[Page]
    label: str = "page"  # "page" for PDFs, "part" for other files and web pages
    text: str = field(init=False, default="")

    def __post_init__(self) -> None:
        self.text = "\n\n".join(f"[{self.label.upper()} {p.number}]\n{p.text}" for p in self.pages)

    def page(self, number: int) -> Page | None:
        for p in self.pages:
            if p.number == number:
                return p
        return None

    @property
    def chars(self) -> int:
        return sum(len(p.text) for p in self.pages)


@dataclass
class Extracted:
    name: str
    pages: list[Page]
    label: str
    links: list[LinkRef]
    truncated: bool = False


# ---------------------------------------------------------------- text tools

_DOT_LEADER = re.compile(r"\.{4,}")
_CURLY = str.maketrans(
    {"\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-", "\u00ad": ""}
)


def squash(text: str) -> str:
    """Normalise text so quotes match despite line breaks, curly quotes, etc."""
    text = unicodedata.normalize("NFKC", text).translate(_CURLY)
    return re.sub(r"\s+", "", text).lower()


def clean_page_text(raw: str) -> str:
    text = _DOT_LEADER.sub(" ", raw)
    lines = [ln.strip() for ln in text.splitlines()]
    out: list[str] = []
    for ln in lines:
        if ln or (out and out[-1]):
            out.append(ln)
    return "\n".join(out).strip()


def _printed_number(raw: str) -> str | None:
    for ln in raw.splitlines()[:6]:
        if re.fullmatch(r"\s*\d{1,3}\s*", ln):
            return ln.strip()
    return None


def split_into_parts(text: str, size: int = 3000) -> list[Page]:
    """Split text into numbered parts so answers can point somewhere."""
    text = text.strip()
    parts: list[Page] = []
    n = 1
    while text:
        if len(text) <= size:
            chunk, text = text, ""
        else:
            cut = text.rfind("\n", 0, size)
            if cut < size * 0.5:
                cut = text.rfind(". ", 0, size)
            if cut < size * 0.5:
                cut = size
            chunk, text = text[:cut], text[cut:]
        chunk = chunk.strip()
        if chunk:
            parts.append(Page(n, chunk, local=n))
            n += 1
    return parts


# --------------------------------------------------------------- link finding

_URL = re.compile(r"https?://[^\s<>\"'\]\[)]+", re.I)
_TRAILING = ".,;:!?)}>'\u201d\u2019"


def clean_url(url: str) -> str:
    return url.strip().rstrip(_TRAILING)


def urls_in_text(text: str) -> list[str]:
    return [clean_url(m.group(0)) for m in _URL.finditer(text)]


def add_links(found: dict[str, LinkRef], urls: list[str], page: int | None) -> None:
    for u in urls:
        u = clean_url(u)
        if u.lower().startswith(("http://", "https://")) and len(u) <= 2000 and u not in found:
            found[u] = LinkRef(u, page)


# ------------------------------------------------------------- file readers

def _read_pdf(data: bytes, name: str) -> Extracted:
    pages: list[Page] = []
    links: dict[str, LinkRef] = {}
    total = 0
    truncated = False
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception:
        raise DocumentError("I couldn't open that PDF. It may be damaged or password protected.")
    with doc:
        if doc.needs_pass:
            raise DocumentError("That PDF is password protected. Remove the password and upload it again.")
        for i, page in enumerate(doc, start=1):
            if i > MAX_PDF_PAGES or total > MAX_TEXT_CHARS:
                truncated = True
                break
            raw = page.get_text()
            text = clean_page_text(raw)
            total += len(text)
            pages.append(Page(i, text, _printed_number(raw), local=i))
            add_links(links, [l["uri"] for l in page.get_links() if l.get("uri")], i)
            add_links(links, urls_in_text(raw), i)
    if not any(len(p.text) > 20 for p in pages):
        raise DocumentError("I can't find any text in that PDF. It may be a scan or photo, which I can't read yet. Try a version you can select text in.")
    return Extracted(name, pages, "page", list(links.values()), truncated)


def _read_docx(data: bytes, name: str) -> Extracted:
    try:
        import docx  # python-docx
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
    except ImportError:  # pragma: no cover
        raise DocumentError("Word files aren't supported on this server.")
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception:
        raise DocumentError("I couldn't open that Word file. Try saving it as a PDF and uploading that.")
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = clean_page_text("\n".join(lines))
    if len(text) < 20:
        raise DocumentError("I can't find any text in that Word file.")
    truncated = len(text) > MAX_TEXT_CHARS
    text = text[:MAX_TEXT_CHARS]
    links: dict[str, LinkRef] = {}
    for rel in d.part.rels.values():
        if rel.reltype == RT.HYPERLINK and rel.is_external:
            add_links(links, [rel.target_ref], None)
    add_links(links, urls_in_text(text), None)
    return Extracted(name, split_into_parts(text), "part", list(links.values()), truncated)


def _read_text(data: bytes, name: str) -> Extracted:
    text = clean_page_text(data.decode("utf-8", errors="replace"))
    if len(text) < 20:
        raise DocumentError("That file has no text in it.")
    truncated = len(text) > MAX_TEXT_CHARS
    text = text[:MAX_TEXT_CHARS]
    links: dict[str, LinkRef] = {}
    add_links(links, urls_in_text(text), None)
    return Extracted(name, split_into_parts(text), "part", list(links.values()), truncated)


def read_upload(filename: str, data: bytes) -> Extracted:
    """Turn an uploaded file into pages of text plus the links found inside it."""
    if not data:
        raise DocumentError("That file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError(f"That file is too big. The limit is {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
    name = Path(filename or "document").name[:120] or "document"
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if data[:5] == b"%PDF-" or ext == "pdf":
        return _read_pdf(data, name)
    if zipfile.is_zipfile(io.BytesIO(data)) and ext in ("docx", ""):
        return _read_docx(data, name)
    if ext in ("txt", "md", "text", ""):
        return _read_text(data, name)
    raise DocumentError("I can read PDF, Word (.docx) and text files. Please upload one of those.")


def load_sample() -> Extracted:
    return read_upload("Sample: planning agreement (155 Mitchell Road).pdf", SAMPLE_PDF.read_bytes())


def to_source(ex: Extracted) -> Source:
    return Source(name=ex.name, kind="main", pages=ex.pages, label=ex.label)


# ------------------------------------------------------- quote verification

def find_quote(source: Source, quote: str, prefer_page: int | None = None) -> int | None:
    """Return the page number where `quote` appears verbatim, else None."""
    if not quote:
        return None
    if "..." in quote or "\u2026" in quote:
        return None  # quotes must be continuous text, not stitched together
    q = squash(quote)
    if len(q) < MIN_QUOTE_CHARS:
        return None
    if prefer_page is not None:
        page = source.page(prefer_page)
        if page and q in squash(page.text):
            return page.number
    for page in source.pages:
        if q in squash(page.text):
            return page.number
    return None


_NUMBER = re.compile(r"\d[\d,]*(?:\.\d+)?")
_REF_WORDS = re.compile(
    r"\b(?:page|pages|clause|clauses|schedule|annexure|item|section|paragraph|part|article|rule|regulation)\s+\d[\d.,]*",
    re.I,
)


def numbers_in(text: str) -> list[str]:
    text = _REF_WORDS.sub(" ", text)
    found = []
    for m in _NUMBER.finditer(text):
        token = m.group(0).replace(",", "").rstrip(".")
        if token:
            found.append(token)
    return found


def numbers_supported(answer: str, evidence: str) -> tuple[bool, str | None]:
    """Every figure in the answer must appear in the evidence text."""
    ev = squash(evidence).replace(",", "")
    for token in numbers_in(answer):
        if token not in ev:
            return False, token
    return True, None


# ------------------------------------------------------------------- config

def load_config() -> dict:
    return json.loads((DATA / "config.json").read_text(encoding="utf-8"))
