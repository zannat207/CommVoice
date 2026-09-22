import io

import pymupdf
import pytest

import documents
from documents import DocumentError, clean_url, read_upload


def make_pdf(pages, link=None, encrypt=False):
    doc = pymupdf.open()
    for i, text in enumerate(pages, start=1):
        page = doc.new_page()
        y = 72
        for line in text.split("\n"):
            page.insert_text((72, y), line)
            y += 16
        if link and i == 1:
            page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(72, 60, 300, 80), "uri": link})
    if encrypt:
        return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    return doc.tobytes()


def test_clean_url_drops_trailing_punctuation():
    assert clean_url("https://example.org/act.") == "https://example.org/act"
    assert clean_url("https://example.org/act),") == "https://example.org/act"


def test_pdf_text_pages_and_links_from_annotations_and_text():
    data = make_pdf(
        ["The tenant must pay rent on the first day of each month.\nSee https://www.legislation.nsw.gov.au/act-1979.", "Page two has more words about the lease and its duties."],
        link="https://example.org/linked",
    )
    ex = read_upload("lease.pdf", data)
    assert ex.label == "page" and [p.number for p in ex.pages] == [1, 2]
    urls = {l.url: l.page for l in ex.links}
    assert urls == {"https://example.org/linked": 1, "https://www.legislation.nsw.gov.au/act-1979": 1}


def test_scanned_pdf_gets_a_plain_message():
    doc = pymupdf.open()
    doc.new_page()
    with pytest.raises(DocumentError) as e:
        read_upload("scan.pdf", doc.tobytes())
    assert "scan" in str(e.value)


def test_password_protected_pdf_is_explained():
    with pytest.raises(DocumentError) as e:
        read_upload("locked.pdf", make_pdf(["Some words here that are long enough."], encrypt=True))
    assert "password" in str(e.value)


def test_docx_text_tables_and_hyperlinks():
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    d = docx.Document()
    d.add_paragraph("This lease starts on the first of March and lasts for one year.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Bond"
    t.rows[0].cells[1].text = "Four weeks rent"
    d.part.relate_to("https://example.org/tenancy-act", RT.HYPERLINK, is_external=True)
    buf = io.BytesIO()
    d.save(buf)
    ex = read_upload("lease.docx", buf.getvalue())
    text = " ".join(p.text for p in ex.pages)
    assert ex.label == "part" and "first of March" in text and "Bond | Four weeks rent" in text
    assert [l.url for l in ex.links] == ["https://example.org/tenancy-act"]


def test_plain_text_file_with_link():
    ex = read_upload("notes.txt", b"You may read the rules at https://example.org/rules, then sign below.")
    assert ex.label == "part" and ex.links[0].url == "https://example.org/rules"


def test_unsupported_and_empty_and_huge_files():
    with pytest.raises(DocumentError):
        read_upload("program.exe", b"\x00\x01\x02\x03" * 10)
    with pytest.raises(DocumentError):
        read_upload("empty.pdf", b"")
    with pytest.raises(DocumentError):
        read_upload("big.txt", b"x" * (documents.MAX_UPLOAD_BYTES + 1))


def test_very_long_text_is_cut_and_flagged():
    ex = read_upload("long.txt", (b"This is a sentence in a very long document. " * 10000))
    assert ex.truncated and sum(len(p.text) for p in ex.pages) <= documents.MAX_TEXT_CHARS + 3200


def test_sample_document_loads_and_has_no_scan_problem():
    ex = documents.load_sample()
    assert len(ex.pages) == 44
